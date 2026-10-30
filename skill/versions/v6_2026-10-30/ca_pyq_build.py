# -*- coding: utf-8 -*-
"""CA-PYQ Match Log builder (v6).

Usage:
  python3 ca_pyq_build.py --logs day1.csv day2.csv ... --today TODAY.json \
      --out-xlsx "CA-PYQ Match Log.xlsx" --out-day-csv today_upload.csv

--logs     = the per-day logs exported from Drive as CSV (11-column schema below; header row required;
             day-sheets whose ninth header still reads 'Representative PYQ (excerpt)' are accepted).
TODAY.json = list of objects: date, topic, source, gs_area, years, papers, pyq_excerpt, match,
             note, prelims [list of str], mains [list of str, each starting with a [Tag]].
             pyq_excerpt holds every matched question, one per line; note holds the three
             tagged lines [Prelims 2027] / [Mains 2027] / [Link].
Writes the formatted workbook (all days, locked design: Calibri 12 everywhere) and a Drive-safe CSV
of today's rows only. This script is the single source of truth for the look — never hand-format.
"""
import argparse, csv, json, math, re
from collections import Counter, OrderedDict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

HEAD = ['Date', 'GS Area', 'Match', 'CA Topic / Headline', 'Prelims Recall (CSE 2027)',
        'Mains Recall (CSE 2027)', 'PYQ Years (all occurrences)', 'Paper(s)',
        'PYQ Excerpts (all, Prelims + Mains)', 'Relevance & Revision Note', 'Source']
FONT, SZ = 'Calibri', 12
NAVY, SLATE = '1F3864', '44546A'
MATCH_FILL = {'Direct': 'C6EFCE', 'Thematic': 'FFF2CC', 'Emerging': 'E7E6E6'}
MATCH_INK = {'Direct': '145A32', 'Thematic': '7D6608', 'Emerging': '3B3838'}
BAND = ['FFFFFF', 'F2F6FB']
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical='top', wrap_text=True)
TOPC = Alignment(vertical='top', horizontal='center', wrap_text=True)

def F(**kw):
    return Font(name=FONT, size=SZ, **kw)

def fill(hex_):
    return PatternFill('solid', fgColor=hex_)

def bullets(xs):
    return '\n'.join('• ' + x.strip() for x in xs if x and x.strip())

def est_height(text, width):
    """Row height (pt) for wrapped Calibri-12 text in a column of given width."""
    cpl = max(8, int(width * 1.15))
    n = sum(max(1, math.ceil(len(seg) / cpl)) for seg in str(text or '').split('\n'))
    return min(409, n * 15.6 + 6)

def mains_rich(text):
    """Bold the [Tag] at the start of each Mains bullet."""
    rt = CellRichText()
    lines = str(text or '').split('\n')
    for i, line in enumerate(lines):
        m = re.match(r'^(•\s*)?(\[[^\]]+\])(.*)$', line)
        if m:
            rt.append(TextBlock(InlineFont(rFont=FONT, sz=SZ), m.group(1) or ''))
            rt.append(TextBlock(InlineFont(rFont=FONT, sz=SZ, b=True, color=NAVY), m.group(2)))
            rt.append(TextBlock(InlineFont(rFont=FONT, sz=SZ), m.group(3) + ('\n' if i < len(lines) - 1 else '')))
        else:
            rt.append(TextBlock(InlineFont(rFont=FONT, sz=SZ), line + ('\n' if i < len(lines) - 1 else '')))
    return rt

def load(log_paths, today_path):
    rows = []
    for log_path in log_paths or []:
        with open(log_path, newline='', encoding='utf-8') as f:
            r = list(csv.reader(f))
        hdr = r[0]
        old = [h.replace('Representative PYQ (excerpt)', 'PYQ Excerpts (all, Prelims + Mains)') for h in hdr]
        assert old[:len(HEAD)] == HEAD, f'Unexpected header in {log_path}: {hdr}'
        rows += [(x + [''] * 11)[:11] for x in r[1:] if x and x[0].strip()]
    if today_path:
        for o in json.load(open(today_path, encoding='utf-8')):
            rows.append([o['date'], o['gs_area'], o['match'], o['topic'], bullets(o['prelims']),
                         bullets(o['mains']), o['years'], o['papers'], o['pyq_excerpt'], o['note'], o['source']])
    seen, out = set(), []
    for x in rows:                                  # de-duplicate on date + topic start
        k = (x[0], x[3][:80])
        if k in seen:
            continue
        seen.add(k); out.append(x)
    order = {id(x): i for i, x in enumerate(out)}
    out.sort(key=lambda x: (x[0], -order[id(x)]), reverse=True)   # newest date first, original order within a date
    return out

def nice_date(d):
    try:
        return datetime.strptime(d, '%Y-%m-%d').strftime('%d %b %Y')
    except ValueError:
        return d

def build(rows, out_xlsx):
    wb = Workbook()
    by_date = OrderedDict()
    for x in rows:
        by_date.setdefault(x[0], []).append(x)

    # ---------- Sheet 1: Revision Cards (phone-friendly, 2 columns) ----------
    ws = wb.active; ws.title = 'Revision Cards'
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 100
    ws.sheet_view.showGridLines = False
    r = 1
    for d, items in by_date.items():
        mix = Counter(x[2] for x in items)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, f"{nice_date(d)}   ·   {len(items)} items   ·   Direct {mix.get('Direct',0)}  |  Thematic {mix.get('Thematic',0)}  |  Emerging {mix.get('Emerging',0)}")
        c.font = F(bold=True, color='FFFFFF'); c.fill = fill(NAVY)
        c.alignment = Alignment(vertical='center', horizontal='left', indent=1)
        ws.row_dimensions[r].height = 26
        r += 1
        for n, x in enumerate(items, 1):
            date, area, match, topic, pre, mai, yrs, pap, exc, note, src = x
            card = [
                (f'#{n}  ·  {match}', topic, 'head'),
                ('Area · Source', f'{area}   ·   {src}', 'meta'),
                ('PYQ years', f'{yrs}   —   {pap}', 'years'),
                ('Prelims recall', pre, 'pre'),
                ('Mains recall', mai, 'mains'),
                ('PYQ excerpts', exc, 'pyq'),
                ('Expected 2027 Qs', note, 'note'),
            ]
            for label, val, kind in card:
                a, b = ws.cell(r, 1, label), ws.cell(r, 2)
                a.font = F(bold=True, color=SLATE); a.alignment = TOP
                a.border = Border(bottom=THIN); b.border = Border(bottom=THIN)
                b.alignment = TOP
                if kind == 'head':
                    a.fill = fill(MATCH_FILL.get(match, 'E7E6E6'))
                    a.font = F(bold=True, color=MATCH_INK.get(match, '3B3838'))
                    b.value = val; b.font = F(bold=True); b.fill = fill('EEF3FA')
                elif kind == 'mains':
                    b.value = mains_rich(val); b.font = F()
                elif kind == 'years':
                    b.value = val; b.font = F(bold=True, color=NAVY)
                elif kind == 'pyq':
                    b.value = val; b.font = F(italic=True, color='404040')
                elif kind == 'pre':
                    b.value = val; b.font = F(); b.fill = fill('F4F9F1')
                else:
                    b.value = val; b.font = F()
                ws.row_dimensions[r].height = max(est_height(val, 100), est_height(label, 20))
                r += 1
            ws.row_dimensions[r].height = 10     # spacer between cards
            r += 1
        r += 1
    ws.freeze_panes = 'A1'
    ws.sheet_properties.tabColor = NAVY
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------- Sheet 2: Master Log (filterable table) ----------
    wl = wb.create_sheet('Master Log')
    widths = {'A': 12, 'B': 18, 'C': 11, 'D': 55, 'E': 55, 'F': 58, 'G': 22, 'H': 16, 'I': 62, 'J': 55, 'K': 18}
    for col, w in widths.items():
        wl.column_dimensions[col].width = w
    wl.append(HEAD)
    for c in wl[1]:
        c.font = F(bold=True, color='FFFFFF'); c.fill = fill(NAVY); c.border = BOX
        c.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
    wl.row_dimensions[1].height = 36
    band = {d: BAND[i % 2] for i, d in enumerate(by_date)}
    prev = None
    for x in rows:
        wl.append(x[:5] + [None] + x[6:])
        i = wl.max_row
        wl.cell(i, 6).value = mains_rich(x[5])
        for j in range(1, 12):
            c = wl.cell(i, j)
            c.font = F(); c.alignment = TOP; c.border = BOX; c.fill = fill(band[x[0]])
        for j in (1, 2, 3, 7, 8):
            wl.cell(i, j).alignment = TOPC
        wl.cell(i, 7).font = F(bold=True, color=NAVY)
        wl.cell(i, 9).font = F(italic=True, color='404040')
        m = wl.cell(i, 3)
        if x[2] in MATCH_FILL:
            m.fill = fill(MATCH_FILL[x[2]]); m.font = F(bold=True, color=MATCH_INK[x[2]])
        if prev is not None and x[0] != prev:
            for j in range(1, 12):
                wl.cell(i, j).border = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style='medium', color=NAVY))
        wl.cell(i, 1).font = F(bold=(x[0] != prev))
        prev = x[0]
        wl.row_dimensions[i].height = max(est_height(v, widths[chr(65 + k)]) for k, v in enumerate(x))
    wl.freeze_panes = 'D2'
    wl.auto_filter.ref = f'A1:K{wl.max_row}'
    wl.sheet_properties.tabColor = '2E75B6'
    wl.print_title_rows = '1:1'
    wl.page_setup.orientation = 'landscape'
    wl.page_setup.fitToWidth = 1; wl.page_setup.fitToHeight = 0
    wl.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------- Sheet 3: PYQ Year Frequency ----------
    wf = wb.create_sheet('PYQ Year Frequency')
    yc, yd = Counter(), Counter()
    for x in rows:
        for y in set(re.findall(r'\b(20[12]\d)\b', x[6] or '')):
            yc[y] += 1
            if x[2] == 'Direct':
                yd[y] += 1
    wf.append(['PYQ Year', 'CA items linked', 'of which Direct', 'Bar'])
    for c in wf[1]:
        c.font = F(bold=True, color='FFFFFF'); c.fill = fill(NAVY); c.border = BOX
        c.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
    mx = max(yc.values()) if yc else 1
    for y in sorted(yc):
        wf.append([y, yc[y], yd[y], '█' * max(1, round(yc[y] * 40 / mx))])
        i = wf.max_row
        for j in range(1, 5):
            c = wf.cell(i, j); c.font = F(); c.border = BOX
            c.alignment = Alignment(horizontal='center' if j < 4 else 'left', vertical='center')
        wf.cell(i, 4).font = F(color='2E75B6')
    wf.append([]); wf.append(['Match mix', 'Items'])
    hr = wf.max_row
    for j in (1, 2):
        wf.cell(hr, j).font = F(bold=True, color='FFFFFF'); wf.cell(hr, j).fill = fill(NAVY)
    for k in ('Direct', 'Thematic', 'Emerging'):
        wf.append([k, sum(1 for x in rows if x[2] == k)])
        i = wf.max_row
        wf.cell(i, 1).fill = fill(MATCH_FILL[k]); wf.cell(i, 1).font = F(bold=True, color=MATCH_INK[k])
        wf.cell(i, 2).font = F()
    for col, w in {'A': 14, 'B': 18, 'C': 16, 'D': 48}.items():
        wf.column_dimensions[col].width = w
    wf.freeze_panes = 'A2'
    wf.sheet_view.showGridLines = False
    wf.sheet_properties.tabColor = SLATE

    # ---------- Sheet 4: How to Use ----------
    wh = wb.create_sheet('How to Use')
    guide = [
        ('What this is', ''),
        ('', "A running log matching each day's current affairs to every UPSC CSE PYQ year in which the same theme was asked (Prelims 2011-2026 and Mains GS1-GS4 2013-2026 from your GitHub PYQ corpus). Each item carries Prelims and Mains key points to recall for CSE 2027."),
        ('Tabs', ''),
        ('Revision Cards', 'Read-first view. One card per news item, newest day first; two columns so it reads on a phone without sideways scrolling.'),
        ('Master Log', 'The same data as a filterable table. Filter Match = Direct or a GS Area; sort by PYQ Years.'),
        ('PYQ Year Frequency', 'How many logged items link to each PYQ year, and how many of those are Direct.'),
        ('Card fields', ''),
        ('#n · Match', 'Direct = a PYQ asked this exact theme. Thematic = the item is a fresh example for a related PYQ. Emerging = not asked yet; high-probability for 2027.'),
        ('PYQ years', 'Every year the theme appeared, Prelims and Mains, with the papers.'),
        ('Prelims recall', 'Testable one-line facts: Acts and years, Articles, bodies, places, numbers.'),
        ('Mains recall', 'Answer-ready points, tagged [Intro data], [Issue], [Judgment], [Committee], [Way forward], [Counterpoint], [Ethics].'),
        ('PYQ excerpts', 'Every past question matched to the item, one line per year: P = Prelims with year, GS1-GS4 = Mains with year and marks. Four years asked means four quoted questions here.'),
        ('Expected 2027 Qs', 'How the examiner would use this item in 2027: one probable Prelims stem, one probable Mains question with paper and marks, and the PYQ years it extends. The two lines test different things by design.'),
        ('Revision routine', ''),
        ('Daily (10 min)', "Read today's cards: Prelims recall aloud, one Mains point per card in your own words."),
        ('Weekly', 'Master Log, filter Match = Direct: write intros using the [Intro data] points; attempt one [Mains 2027] question in full.'),
        ('Before a test', 'Filter by GS Area; revise the Prelims recall column and the [Prelims 2027] lines only.'),
        ('Daily run', 'Upload or drop the day\'s paper in UPSC Prep/Daily Newspaper/ and ask for the CA-PYQ match (skill: ca-pyq-daily-match).'),
    ]
    wh.append(['Field', 'Explanation'])
    for c in wh[1]:
        c.font = F(bold=True, color='FFFFFF'); c.fill = fill(NAVY)
    for a, b in guide:
        wh.append([a, b])
        i = wh.max_row
        wh.cell(i, 1).font = F(bold=True); wh.cell(i, 2).font = F()
        wh.cell(i, 1).alignment = TOP; wh.cell(i, 2).alignment = TOP
        if a and not b:
            for j in (1, 2):
                wh.cell(i, j).fill = fill('D9E2F3')
        wh.row_dimensions[i].height = est_height(b, 100)
    wh.column_dimensions['A'].width = 22; wh.column_dimensions['B'].width = 100
    wh.freeze_panes = 'A2'
    wh.sheet_view.showGridLines = False
    wh.sheet_properties.tabColor = 'A6A6A6'
    wb.save(out_xlsx)

def sheet_safe(v):
    """Google Sheets strips one leading apostrophe and evaluates =,+,- as formulas: guard them."""
    v = str(v)
    return "'" + v if v[:1] in ("'", '=', '+', '-') else v

def write_csv(rows, path):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(HEAD); w.writerows([[sheet_safe(v) for v in x] for x in rows])

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--logs', nargs='*', default=[]); ap.add_argument('--today')
    ap.add_argument('--out-xlsx', required=True); ap.add_argument('--out-day-csv')
    ap.add_argument('--monthly-dir', help='also write one workbook per month: <dir>/CA-PYQ Match Log YYYY-MM.xlsx')
    a = ap.parse_args()
    rows = load(a.logs, a.today)
    build(rows, a.out_xlsx)
    if a.monthly_dir:
        import os
        os.makedirs(a.monthly_dir, exist_ok=True)
        for ym in sorted({x[0][:7] for x in rows}):
            month_rows = [x for x in rows if x[0][:7] == ym]
            build(month_rows, os.path.join(a.monthly_dir, f'CA-PYQ Match Log {ym}.xlsx'))
            print('monthly:', ym, len(month_rows), 'rows')
    if a.today and a.out_day_csv:
        d = json.load(open(a.today, encoding='utf-8'))[0]['date']
        day = [x for x in rows if x[0] == d]
        write_csv(day, a.out_day_csv)
        chk = list(csv.reader(open(a.out_day_csv, newline='', encoding='utf-8')))
        assert len(chk) == len(day) + 1 and all(len(x) == 11 for x in chk), 'CSV round-trip failed'
        print('today CSV:', a.out_day_csv, len(day), 'rows,', len(open(a.out_day_csv, encoding='utf-8').read()), 'chars')
    print('workbook rows:', len(rows), '| by date:', sorted(Counter(x[0] for x in rows).items(), reverse=True),
          '| match:', dict(Counter(x[2] for x in rows)))
